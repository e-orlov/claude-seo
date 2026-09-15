from __future__ import annotations

import copy
import datetime as dt
import json
import os
import stat
import subprocess
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

import yaml
import openpyxl


REPO_ROOT = Path(__file__).resolve().parents[3]
SCRIPT_DIR = REPO_ROOT / ".claude" / "skills" / "seo-geo-audit" / "scripts"
sys.path.insert(0, str(SCRIPT_DIR))

from estimate_sistrix_cost import build_plan  # noqa: E402
from extract_zip import extract  # noqa: E402
from inspect_mcp_probe import parse_response, shape, xml_path  # noqa: E402
from preflight_common import canonical_sha256, load_manifest  # noqa: E402
from source_adapters import classify_headers, inspect_sentiment_mhtml, inspect_tabular  # noqa: E402


SF_CODES = ["SF", "RAW", "REN", "PSI", "GSC-SA", "GSC-UI"]
SX_CODES = ["SX-M", "SX-O", "SX-C", "SX-P", "SX-PC", "SX-S"]
SX_ENDPOINTS = {
    "SX-M": "ai.models",
    "SX-O": "ai.check.overview",
    "SX-C": "ai.check.competitors",
    "SX-P": "ai.check.prompts",
    "SX-PC": "ai.check.prompts.count",
    "SX-S": "ai.check.sources",
}


def now() -> str:
    return dt.datetime.now(dt.timezone.utc).replace(microsecond=0).isoformat().replace("+00:00", "Z")


def run_script(name: str, *arguments: str, expected: int = 0) -> subprocess.CompletedProcess[str]:
    result = subprocess.run(
        [sys.executable, str(SCRIPT_DIR / name), *arguments],
        cwd=REPO_ROOT,
        text=True,
        capture_output=True,
        check=False,
    )
    if result.returncode != expected:
        raise AssertionError(f"{name} returned {result.returncode}, expected {expected}\nstdout={result.stdout}\nstderr={result.stderr}")
    return result


def write_json(path: Path, value: object) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def valid_config(root: Path) -> dict:
    config = json.loads((REPO_ROOT / ".claude" / "geo-audit" / "fixtures" / "minimal-complete" / "audit-config.json").read_text(encoding="utf-8"))
    config["client_id"] = "example-com"
    config["brand"] = "Example"
    config["domain"] = "example.com"
    config["country"] = "de"
    config["language"] = "de"
    config["sistrix"]["country"] = "de"
    config["sistrix"]["models"] = ["chatgpt", "aio"]
    inputs = root / "inputs"
    inputs.mkdir(parents=True, exist_ok=True)
    paths = {
        "ahrefs_backlinks": inputs / "backlinks.csv",
        "ahrefs_refdomains": inputs / "refdomains.csv",
        "ahrefs_broken_backlinks": inputs / "broken-backlinks.csv",
        "dejan": inputs / "dejan.csv",
        "gsc_generative_ai_export": inputs / "gsc-generative.csv",
    }
    paths["ahrefs_backlinks"].write_text("Referring page URL,Target URL\nhttps://ref.test/a,https://example.com/a\n", encoding="utf-8")
    paths["ahrefs_refdomains"].write_text("Referring domain,Backlinks\nref.test,2\n", encoding="utf-8")
    paths["ahrefs_broken_backlinks"].write_text("Referring page URL,Broken target URL,HTTP code\nhttps://ref.test/a,https://example.com/missing,404\n", encoding="utf-8")
    paths["dejan"].write_text("Agent,Access\nOAI-SearchBot,Allowed for example.com\n", encoding="utf-8")
    paths["gsc_generative_ai_export"].write_text("Date,Page,Impressions\n2026-09-14,https://example.com/a,12\n", encoding="utf-8")
    for key, path in paths.items():
        config["inputs"][key] = str(path)
    mhtml = inputs / "example-sentiment.mhtml"
    mhtml.write_bytes(synthetic_mhtml())
    config["inputs"]["sistrix_sentiment_mhtml"] = [str(mhtml)]
    return config


def synthetic_mhtml() -> bytes:
    return (
        "From: <Saved by Blink>\r\n"
        "Subject: Example: Sentiment - SISTRIX for AI\r\n"
        "MIME-Version: 1.0\r\n"
        "Content-Type: multipart/related; boundary=fixture\r\n"
        "Snapshot-Content-Location: https://ai.sistrix.com/research/example/sentiment\r\n\r\n"
        "--fixture\r\n"
        "Content-Type: text/html; charset=utf-8\r\n"
        "Content-Transfer-Encoding: 8bit\r\n"
        "Content-Location: https://ai.sistrix.com/research/example/sentiment\r\n\r\n"
        "<html><head><title>Example: Sentiment - SISTRIX for AI</title></head>"
        "<body><h1>Sentiment-Score +42</h1><p>Example example.com</p>"
        "<p>Lob 71% Kritik 29% ChatGPT AI Overview</p></body></html>\r\n"
        "--fixture--\r\n"
    ).encode("utf-8")


def create_probe(root: Path, server: str, partial_code: str | None = None) -> Path:
    codes = SF_CODES if server == "screaming_frog" else SX_CODES
    response_path = root / f"{server}-response.json"
    response = {
        "domain": "example.com",
        "country": "de" if server == "sistrix" else None,
        "snapshot_at": now(),
        "counts": {code: 1 for code in codes},
        "models": ["chatgpt", "aio"],
    }
    write_json(response_path, response)
    calls = [{
        "call_id": f"{server}-call",
        "callable_name": f"fixture_{server}_read",
        "request_parameters": {"domain": "example.com"},
        "response_path": str(response_path),
        "response_sha256": "0" * 64,
        "response_schema_fingerprint": "0" * 64,
        "recorded_at": now(),
    }]
    results = []
    for code in codes:
        results.append({
            "source_code": code,
            "call_ids": [f"{server}-call"],
            "access_status": "pass",
            "data_status": "partial" if code == partial_code else "present",
            "scope_status": "match",
            "freshness_status": "current",
            "snapshot_at": now(),
            "record_count": 1,
            "valid_empty": False,
            "semantic_checks": ["draft"],
            "assertions": [
                {
                    "call_id": f"{server}-call", "proof_role": "scope_domain", "locator_type": "json_pointer", "locator": "/domain",
                    "operator": "equals", "expected": "example.com", "description": "domain matches",
                },
                {
                    "call_id": f"{server}-call", "proof_role": "record_count", "locator_type": "json_pointer", "locator": f"/counts/{code}",
                    "operator": "equals", "expected": 1, "description": f"{code} row count is present",
                },
                {
                    "call_id": f"{server}-call", "proof_role": "snapshot", "locator_type": "json_pointer", "locator": "/snapshot_at",
                    "operator": "equals", "expected": response["snapshot_at"], "description": "snapshot is anchored",
                },
                {
                    "call_id": f"{server}-call", "proof_role": "schema", "locator_type": "json_pointer", "locator": f"/counts/{code}",
                    "operator": "exists", "expected": None, "description": "source field exists",
                },
            ],
            "details": {"endpoint": SX_ENDPOINTS[code]} if server == "sistrix" else {},
            "limitations": ["fixture record-level gap"] if code == partial_code else [],
        })
        if server == "sistrix":
            results[-1]["assertions"].append({
                "call_id": f"{server}-call", "proof_role": "scope_country", "locator_type": "json_pointer", "locator": "/country",
                "operator": "equals", "expected": "de", "description": "country matches",
            })
            if code == "SX-M":
                for model in ("chatgpt", "aio"):
                    results[-1]["assertions"].append({
                        "call_id": f"{server}-call", "proof_role": "scope_model", "locator_type": "json_pointer", "locator": "/models",
                        "operator": "contains", "expected": model, "description": f"model {model} is available",
                    })
    draft = {
        "package_type": "geo_mcp_probe",
        "schema_version": 1,
        "server": server,
        "adapter_version": 1,
        "captured_at": now(),
        "scope": {"domain": "example.com", "country": "de" if server == "sistrix" else None, "models": ["chatgpt", "aio"] if server == "sistrix" else []},
        "calls": calls,
        "source_results": results,
        "cost_plan": build_plan("mcp", 2, None, {}) if server == "sistrix" else None,
        "verification": {"passed": True, "verified_at": now(), "checks_run": 1, "failures": []},
        "package_sha256": "0" * 64,
    }
    draft_path = root / f"{server}-draft.json"
    output_path = root / f"{server}-probe.json"
    write_json(draft_path, draft)
    arguments = [
        "--draft", str(draft_path), "--output", str(output_path), "--expected-domain", "example.com",
    ]
    if server == "sistrix":
        arguments += ["--expected-country", "de", "--required-models", "chatgpt,aio"]
    run_script("inspect_mcp_probe.py", *arguments)
    return output_path


class PreflightRuntimeTests(unittest.TestCase):
    def test_init_is_resumable_and_revisioned(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            config_path = root / "audit.json"
            write_json(config_path, valid_config(root))
            run_root = root / "runs"
            first = json.loads(run_script("init_run.py", "--config", str(config_path), "--run-root", str(run_root)).stdout)
            resumed = json.loads(run_script("init_run.py", "--config", str(config_path), "--run-root", str(run_root)).stdout)
            revised = json.loads(run_script("init_run.py", "--config", str(config_path), "--run-root", str(run_root), "--new-revision").stdout)
            self.assertEqual(first["run_id"], resumed["run_id"])
            self.assertTrue(first["run_id"].endswith("-r1"))
            self.assertTrue(revised["run_id"].endswith("-r2"))

    def test_missing_config_lists_only_missing_identity(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            result = run_script(
                "init_run.py", "--config", str(root / "audit.yaml"), "--run-root", str(root / "runs"),
                "--domain", "example.com", "--brand", "Example", expected=2,
            )
            output = json.loads(result.stdout)
            self.assertEqual(output["status"], "WAITING_FOR_CONFIG")
            self.assertIn("country", output["missing_fields"])
            self.assertIn("language", output["missing_fields"])
            self.assertNotIn("domain", output["missing_fields"])

    def test_zip_traversal_and_symlink_are_rejected(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            bad = root / "bad.zip"
            with zipfile.ZipFile(bad, "w") as bundle:
                bundle.writestr("../escape.csv", "a,b\n1,2\n")
            with self.assertRaisesRegex(ValueError, "Unsafe ZIP"):
                extract(bad, root / "out")

            symlink = root / "symlink.zip"
            link_info = zipfile.ZipInfo("exports/link.csv")
            link_info.create_system = 3
            link_info.external_attr = (stat.S_IFLNK | 0o777) << 16
            with zipfile.ZipFile(symlink, "w") as bundle:
                bundle.writestr(link_info, "../outside.csv")
            with self.assertRaisesRegex(ValueError, "symlinks are forbidden"):
                extract(symlink, root / "symlink-out")

            good = root / "good.zip"
            with zipfile.ZipFile(good, "w") as bundle:
                bundle.writestr("exports/backlinks.csv", "a,b\n1,2\n")
            manifest = extract(good, root / "safe")
            self.assertEqual(manifest["file_count"], 1)
            self.assertTrue(Path(manifest["files"][0]["path"]).is_file())

    def test_multilingual_tabular_signatures(self) -> None:
        german = ["Verweisende Seite URL", "Ziel URL", "Statuscode"]
        candidates = classify_headers(german, "broken-backlinks.csv")
        self.assertEqual(candidates[0]["source_code"], "AH-BB")
        self.assertTrue(candidates[0]["complete_signature"])
        russian = classify_headers(["Бот", "Доступ"], "dejan.csv")
        self.assertEqual(russian[0]["source_code"], "DJ")
        self.assertTrue(russian[0]["complete_signature"])

    def test_mhtml_mime_decoding_and_no_remote_load(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary) / "sentiment.mhtml"
            path.write_bytes(synthetic_mhtml())
            result = inspect_sentiment_mhtml(path, "Example", "example.com", "visibility_30d")
            self.assertFalse(result["check"]["blocking"])
            self.assertEqual(result["check"]["data_status"], "present")
            self.assertFalse(result["parsed"]["remote_resources_loaded"])
            malformed = path.with_name("not-multipart.mhtml")
            malformed.write_text("<html><body>Sentiment</body></html>", encoding="utf-8")
            failed = inspect_sentiment_mhtml(malformed, "Example", "example.com", "visibility_30d")
            self.assertEqual(failed["check"]["data_status"], "malformed")
            self.assertTrue(failed["check"]["blocking"])

    def test_valid_empty_requires_explicit_attestation(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary) / "backlinks.csv"
            path.write_text("Referring page URL,Target URL\n", encoding="utf-8")
            unproven = inspect_tabular(path, "AH-BL", "example.com", "visibility_30d", valid_empty_attested=False)
            proven = inspect_tabular(path, "AH-BL", "example.com", "visibility_30d", valid_empty_attested=True)
            self.assertTrue(unproven["check"]["blocking"])
            self.assertEqual(unproven["check"]["data_status"], "partial")
            self.assertFalse(proven["check"]["blocking"])
            self.assertEqual(proven["check"]["data_status"], "present")
            self.assertTrue(proven["artifact"]["valid_empty"])

            old = dt.datetime.now().timestamp() - 40 * 24 * 60 * 60
            os.utime(path, (old, old))
            stale = inspect_tabular(path, "AH-BL", "example.com", "visibility_30d", valid_empty_attested=True)
            self.assertEqual(stale["check"]["freshness_status"], "stale")
            self.assertTrue(stale["check"]["blocking"])

    def test_explicit_binding_cannot_override_wrong_target_domain(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary) / "backlinks.csv"
            path.write_text(
                "Referring page URL,Target URL\nhttps://ref.test/a,https://wrong.example/a\n",
                encoding="utf-8",
            )
            result = inspect_tabular(path, "AH-BL", "example.com", "visibility_30d", explicit_binding=True)
            self.assertEqual(result["check"]["scope_status"], "mismatch")
            self.assertEqual(result["check"]["details"]["scope_method"], "content_contradiction")
            self.assertTrue(result["check"]["blocking"])

    def test_stored_dejan_html_is_supported_without_network(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary) / "dejan.html"
            path.write_text("<html><body><h1>Agent access</h1><p>GPTBot: Allowed</p><p>ClaudeBot: Blocked</p></body></html>", encoding="utf-8")
            result = inspect_tabular(path, "DJ", "example.com", "crawl_14d", explicit_binding=True)
            self.assertFalse(result["check"]["blocking"])
            self.assertEqual(result["artifact"]["record_count"], 2)
            self.assertTrue(result["check"]["details"]["html_text_signature"])

    def test_gsc_workbook_combines_page_country_and_device_sheets(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary) / "gsc-generative.xlsx"
            workbook = openpyxl.Workbook()
            page = workbook.active
            page.title = "Pages"
            page.append(["Date", "Page", "Impressions"])
            page.append(["2026-09-14", "https://example.com/a", 12])
            country = workbook.create_sheet("Countries")
            country.append(["Date", "Country", "Impressions"])
            country.append(["2026-09-14", "DE", 12])
            device = workbook.create_sheet("Devices")
            device.append(["Date", "Device", "Impressions"])
            device.append(["2026-09-14", "MOBILE", 12])
            workbook.save(path)
            result = inspect_tabular(path, "GSC-GAI", "example.com", "visibility_30d", explicit_binding=True)
            self.assertFalse(result["check"]["blocking"])
            self.assertEqual(result["artifact"]["record_count"], 3)
            self.assertEqual(set(result["check"]["details"]["semantic_mapping"]), {"date", "page", "country", "device", "impressions"})
            self.assertEqual(len(result["check"]["details"]["sheet_profiles"]), 3)

    def test_cost_plan_keeps_unknown_upper_bound(self) -> None:
        plan = build_plan("mcp", 2, None, {})
        self.assertEqual(plan["total_credit_bounds"], {"lower": 0, "expected": 0, "upper": 0})
        self.assertEqual(plan["total_request_bounds"]["lower"], 10)
        self.assertIsNone(plan["total_request_bounds"]["upper"])

    def test_probe_assertions_are_response_anchored(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            probe = create_probe(root, "sistrix")
            package = json.loads(probe.read_text(encoding="utf-8"))
            self.assertTrue(package["verification"]["passed"])
            self.assertEqual(package["verification"]["checks_run"], 32)
            self.assertEqual(package["package_sha256"], canonical_sha256(package, "package_sha256"))
            second = root / "sistrix-probe-second.json"
            run_script(
                "inspect_mcp_probe.py", "--draft", str(root / "sistrix-draft.json"),
                "--output", str(second), "--expected-domain", "example.com",
                "--expected-country", "de", "--required-models", "chatgpt,aio",
            )
            self.assertEqual(probe.read_bytes(), second.read_bytes())

            response = json.loads((root / "sistrix-response.json").read_text(encoding="utf-8"))
            response["domain"] = "wrong.example"
            write_json(root / "sistrix-response.json", response)
            draft = root / "sistrix-draft.json"
            result = subprocess.run(
                [sys.executable, str(SCRIPT_DIR / "inspect_mcp_probe.py"), "--draft", str(draft), "--output", str(root / "invalid.json"), "--expected-domain", "example.com", "--expected-country", "de"],
                cwd=REPO_ROOT, text=True, capture_output=True, check=False,
            )
            self.assertNotEqual(result.returncode, 0)

    def test_sistrix_xml_shape_and_attributes_are_inspectable(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary) / "models.xml"
            path.write_text(
                '<?xml version="1.0"?><response><method>ai.models</method><answer>'
                '<models llm_label="ChatGPT" llm_model="chatgpt"/>'
                '<models llm_label="AI Overview" llm_model="aio"/>'
                '</answer><credits used="0"/></response>',
                encoding="utf-8",
            )
            kind, parsed = parse_response(path)
            self.assertEqual(kind, "xml")
            self.assertEqual(xml_path(parsed, "./answer/models/@llm_model"), ["chatgpt", "aio"])
            self.assertIn("children", shape(parsed))

    def test_full_preflight_ready_then_detects_source_change(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            config_path = root / "audit.json"
            config = valid_config(root)
            write_json(config_path, config)
            run_root = root / "runs"
            created = json.loads(run_script("init_run.py", "--config", str(config_path), "--run-root", str(run_root)).stdout)
            manifest_path = Path(created["manifest_path"])
            sf_probe = create_probe(root, "screaming_frog")
            sx_probe = create_probe(root, "sistrix")
            arguments = [
                "--manifest", str(manifest_path), "--sf-probe", str(sf_probe), "--sistrix-probe", str(sx_probe),
            ]
            ready = json.loads(run_script("preflight.py", *arguments).stdout)
            self.assertEqual(ready["status"], "READY")
            manifest = load_manifest(manifest_path)
            self.assertEqual(len(manifest["source_checks"]), 18)
            self.assertFalse(any(item["blocking"] for item in manifest["source_checks"]))

            resumed = json.loads(run_script("preflight.py", *arguments).stdout)
            self.assertEqual(resumed["status"], "READY")
            backlink_path = Path(config["inputs"]["ahrefs_backlinks"])
            backlink_path.write_text(backlink_path.read_text(encoding="utf-8") + "https://two.test/b,https://example.com/b\n", encoding="utf-8")
            changed = json.loads(run_script("preflight.py", *arguments, expected=2).stdout)
            self.assertEqual(changed["status"], "BLOCKED")
            self.assertIn("--new-revision", changed["action"])

    def test_incomplete_preflight_does_not_freeze_partial_source_set(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            config_path = root / "audit.json"
            config = valid_config(root)
            Path(config["inputs"]["ahrefs_backlinks"]).unlink()
            write_json(config_path, config)
            created = json.loads(run_script("init_run.py", "--config", str(config_path), "--run-root", str(root / "runs")).stdout)
            result = json.loads(run_script("preflight.py", "--manifest", created["manifest_path"], expected=2).stdout)
            self.assertEqual(result["status"], "WAITING_FOR_INPUTS")
            self.assertIsNone(load_manifest(Path(created["manifest_path"]))["source_set_sha256"])

    def test_ready_with_gaps_requires_present_sources_and_record_gap(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            config_path = root / "audit.json"
            write_json(config_path, valid_config(root))
            created = json.loads(run_script("init_run.py", "--config", str(config_path), "--run-root", str(root / "runs")).stdout)
            sf_probe = create_probe(root, "screaming_frog", partial_code="PSI")
            sx_probe = create_probe(root, "sistrix")
            output = json.loads(run_script(
                "preflight.py", "--manifest", created["manifest_path"],
                "--sf-probe", str(sf_probe), "--sistrix-probe", str(sx_probe),
            ).stdout)
            self.assertEqual(output["status"], "READY_WITH_GAPS")
            manifest = load_manifest(Path(created["manifest_path"]))
            self.assertEqual(manifest["record_gaps"][0]["source_code"], "PSI")
            self.assertFalse(any(item["blocking"] for item in manifest["source_checks"]))


if __name__ == "__main__":
    unittest.main()
