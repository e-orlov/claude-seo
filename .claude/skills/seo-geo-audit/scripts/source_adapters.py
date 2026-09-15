#!/usr/bin/env python3
"""Secure file profiling helpers for mandatory GEO inputs."""

from __future__ import annotations

import csv
import datetime as dt
import hashlib
import io
import json
import re
import unicodedata
import zipfile
from email import policy
from email.parser import BytesParser
from html.parser import HTMLParser
from pathlib import Path
from typing import Any, Iterable

from preflight_common import artifact_id, canonical_sha256, freshness, sha256_file


ADAPTER_VERSION = 1
MAX_TABULAR_BYTES = 512 * 1024 * 1024
MAX_MHTML_BYTES = 128 * 1024 * 1024
SUPPORTED_TABLE_EXTENSIONS = {".csv", ".tsv", ".txt", ".xlsx", ".xlsm", ".html", ".htm"}


def norm(value: str) -> str:
    value = "".join(
        char for char in unicodedata.normalize("NFKD", value.casefold())
        if not unicodedata.combining(char)
    )
    return re.sub(r"[^\w]+", "_", value, flags=re.UNICODE).strip("_")


ALIASES = {
    "referring_url": {
        "referring_page_url", "referring_url", "source_url", "verweisende_seite_url",
        "verweisende_url", "url_der_verweisenden_seite", "ссылающаяся_страница_url",
    },
    "referring_domain": {
        "referring_domain", "referring_domains", "verweisende_domain", "verweisende_domains",
        "ссылающийся_домен",
    },
    "target_url": {
        "target_url", "target", "destination_url", "link_url", "broken_target_url",
        "ziel_url", "zielseite", "url_zielseite", "целевая_страница_url",
    },
    "http_status": {
        "http_code", "status_code", "http_status", "target_http_code", "broken_status",
        "http_statuscode", "statuscode", "код_http",
    },
    "lost_date": {"lost", "lost_date", "link_lost", "verloren", "verlustdatum", "дата_потери"},
    "backlinks": {"backlinks", "backlink_count", "links", "ruckverweise", "обратные_ссылки"},
    "date": {"date", "day", "datum", "дата"},
    "page": {"page", "url", "landing_page", "seite", "страница"},
    "country": {"country", "land", "страна"},
    "device": {"device", "gerat", "geraet", "устройство"},
    "impressions": {"impressions", "ai_impressions", "impressionen", "показы"},
    "agent": {"agent", "crawler", "bot", "user_agent", "useragent", "робот", "бот"},
    "access": {"access", "allowed", "crawlability", "status", "zugriff", "erlaubt", "доступ"},
}


SIGNATURES = {
    "AH-BL": {"all": {"referring_url", "target_url"}, "any": set()},
    "AH-RD": {"all": {"referring_domain"}, "any": set()},
    "AH-BB": {"all": {"referring_url", "target_url"}, "any": {"http_status", "lost_date"}},
    "GSC-GAI": {"all": {"date", "impressions"}, "any": {"page", "country", "device"}},
    "DJ": {"all": {"agent", "access"}, "any": set()},
}


FILENAME_HINTS = {
    "AH-BL": ("backlink", "backlinks"),
    "AH-RD": ("referring domain", "refdomains", "referring_domains"),
    "AH-BB": ("broken backlink", "broken_backlink", "broken-backlink"),
    "GSC-GAI": ("generative", "gen ai", "gen_ai", "gai"),
    "DJ": ("dejan", "agent", "crawlability"),
}


class VisibleTextParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self._skip = 0
        self._in_title = False
        self.text: list[str] = []
        self.title: list[str] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag.casefold() in {"script", "style", "noscript", "template"}:
            self._skip += 1
        if tag.casefold() == "title":
            self._in_title = True

    def handle_endtag(self, tag: str) -> None:
        if tag.casefold() in {"script", "style", "noscript", "template"} and self._skip:
            self._skip -= 1
        if tag.casefold() == "title":
            self._in_title = False

    def handle_data(self, data: str) -> None:
        clean = " ".join(data.split())
        if not clean or self._skip:
            return
        self.text.append(clean)
        if self._in_title:
            self.title.append(clean)


def decode_bytes(payload: bytes) -> tuple[str, str]:
    for encoding in ("utf-8-sig", "utf-16", "cp1252", "latin-1"):
        try:
            text = payload.decode(encoding)
            if "\x00" not in text:
                return text, encoding
        except UnicodeDecodeError:
            continue
    raise ValueError("No safe text decoding succeeded")


def header_mapping(headers: Iterable[str]) -> dict[str, str]:
    normalized = {norm(header): header for header in headers if str(header).strip()}
    mapped: dict[str, str] = {}
    for semantic, aliases in ALIASES.items():
        for alias in aliases:
            if norm(alias) in normalized:
                mapped[semantic] = normalized[norm(alias)]
                break
    return mapped


def signature_match(source_code: str, mapping: dict[str, str]) -> tuple[bool, list[str]]:
    signature = SIGNATURES[source_code]
    present = set(mapping)
    missing = sorted(signature["all"] - present)
    if signature["any"] and not (signature["any"] & present):
        missing.append("one_of:" + "|".join(sorted(signature["any"])))
    return not missing, missing


def classify_headers(headers: list[str], filename: str = "") -> list[dict[str, Any]]:
    mapping = header_mapping(headers)
    name = filename.casefold().replace("-", " ").replace("_", " ")
    candidates = []
    for source_code, signature in SIGNATURES.items():
        present = set(mapping)
        required_count = len(signature["all"]) + (1 if signature["any"] else 0)
        hit_count = len(signature["all"] & present) + (1 if signature["any"] & present else 0)
        complete, missing = signature_match(source_code, mapping)
        filename_hit = any(hint.replace("_", " ") in name for hint in FILENAME_HINTS[source_code])
        if source_code == "AH-BL" and "broken" in name:
            filename_hit = False
        score = hit_count / required_count + (0.2 if filename_hit else 0.0)
        candidates.append({
            "source_code": source_code,
            "score": round(score, 4),
            "complete_signature": complete,
            "missing_semantics": missing,
            "mapping": mapping,
            "filename_hint": filename_hit,
        })
    return sorted(candidates, key=lambda item: (-item["complete_signature"], -item["score"], item["source_code"]))


def _csv_profile(path: Path) -> tuple[list[str], int, list[str], dict[str, Any]]:
    payload = path.read_bytes()
    if len(payload) > MAX_TABULAR_BYTES:
        raise ValueError(f"Tabular input exceeds {MAX_TABULAR_BYTES} bytes")
    text, encoding = decode_bytes(payload)
    sample = text[:65536]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = "\t" if path.suffix.casefold() == ".tsv" else ","
    rows = csv.reader(io.StringIO(text), delimiter=delimiter)
    preview = []
    for _ in range(12):
        try:
            preview.append(next(rows))
        except StopIteration:
            break
    if not preview:
        return [], 0, [], {"encoding": encoding, "delimiter": delimiter, "header_row": None}
    best_index = max(range(len(preview)), key=lambda index: len(header_mapping(preview[index])))
    headers = [str(value).strip() for value in preview[best_index]]
    record_count = max(0, len(preview) - best_index - 1)
    sample_values = [str(value) for row in preview[best_index + 1:] for value in row]
    for row in rows:
        if any(str(value).strip() for value in row):
            record_count += 1
            if len(sample_values) < 2000:
                sample_values.extend(str(value) for value in row)
    return headers, record_count, sample_values, {
        "encoding": encoding, "delimiter": delimiter, "header_row": best_index + 1,
    }


def _xlsx_profile(path: Path) -> tuple[list[str], int, list[str], dict[str, Any]]:
    if path.stat().st_size > MAX_TABULAR_BYTES:
        raise ValueError(f"Workbook exceeds {MAX_TABULAR_BYTES} bytes")
    try:
        import openpyxl
    except ImportError as exc:
        raise ValueError("openpyxl is required for XLSX inputs") from exc
    validate_office_archive(path)
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        all_headers: list[str] = []
        total_records = 0
        sample_values: list[str] = []
        sheet_profiles = []
        for sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
            preview = []
            iterator = sheet.iter_rows(values_only=True)
            for _ in range(12):
                try:
                    preview.append(["" if value is None else str(value) for value in next(iterator)])
                except StopIteration:
                    break
            if not preview:
                continue
            best_index = max(range(len(preview)), key=lambda index: len(header_mapping(preview[index])))
            headers = [value.strip() for value in preview[best_index]]
            if not header_mapping(headers):
                continue
            record_count = max(0, len(preview) - best_index - 1)
            values = [value for row in preview[best_index + 1:] for value in row]
            for row in iterator:
                if any(value not in (None, "") for value in row):
                    record_count += 1
                    if len(values) < 2000:
                        values.extend("" if value is None else str(value) for value in row)
            for header in headers:
                if norm(header) not in {norm(item) for item in all_headers}:
                    all_headers.append(header)
            total_records += record_count
            if len(sample_values) < 2000:
                sample_values.extend(values[:2000 - len(sample_values)])
            sheet_profiles.append({
                "sheet": sheet_name,
                "header_row": best_index + 1,
                "headers": headers,
                "record_count": record_count,
            })
        if not sheet_profiles:
            return [], 0, [], {"encoding": "binary-xlsx", "delimiter": None, "header_row": None, "sheet_profiles": []}
        return all_headers, total_records, sample_values, {
            "encoding": "binary-xlsx", "delimiter": None, "header_row": None, "sheet_profiles": sheet_profiles,
        }
    finally:
        workbook.close()


def _html_table_profile(path: Path) -> tuple[list[str], int, list[str], dict[str, Any]]:
    payload = path.read_bytes()
    if len(payload) > MAX_TABULAR_BYTES:
        raise ValueError(f"HTML table input exceeds {MAX_TABULAR_BYTES} bytes")
    text, encoding = decode_bytes(payload)
    try:
        from lxml import html
    except ImportError as exc:
        raise ValueError("lxml is required for HTML table inputs") from exc
    tree = html.fromstring(text)
    tables = tree.xpath("//table")
    if not tables:
        raise ValueError("No HTML table found")
    best = max(tables, key=lambda table: len(table.xpath(".//tr")))
    rows = [[" ".join(cell.itertext()).strip() for cell in row.xpath("./th|./td")] for row in best.xpath(".//tr")]
    rows = [row for row in rows if any(row)]
    if not rows:
        return [], 0, [], {"encoding": encoding, "delimiter": None, "header_row": None}
    best_index = max(range(min(10, len(rows))), key=lambda index: len(header_mapping(rows[index])))
    headers = rows[best_index]
    data = rows[best_index + 1:]
    return headers, len(data), [value for row in data[:200] for value in row], {
        "encoding": encoding, "delimiter": None, "header_row": best_index + 1,
    }


def _dejan_html_profile(path: Path) -> tuple[list[str], int, list[str], dict[str, Any]]:
    payload = path.read_bytes()
    if len(payload) > MAX_TABULAR_BYTES:
        raise ValueError(f"Dejan HTML input exceeds {MAX_TABULAR_BYTES} bytes")
    text, encoding = decode_bytes(payload)
    parser = VisibleTextParser()
    parser.feed(text)
    visible = "\n".join(parser.text)
    agents = {
        name for name in (
            "GPTBot", "OAI-SearchBot", "ChatGPT-User", "ClaudeBot", "Claude-SearchBot",
            "PerplexityBot", "Perplexity-User", "Google-Extended", "Applebot-Extended",
        ) if name.casefold() in visible.casefold()
    }
    access_terms = ("allowed", "blocked", "allow", "disallow", "erlaubt", "blockiert", "zugriff")
    if not agents or not any(term in visible.casefold() for term in access_terms):
        raise ValueError("Stored Dejan HTML lacks agent/access result signals")
    return ["Agent", "Access"], len(agents), [visible[:20000]], {
        "encoding": encoding, "delimiter": None, "header_row": None, "html_text_signature": True,
    }


def profile_table(path: Path) -> tuple[list[str], int, list[str], dict[str, Any]]:
    suffix = path.suffix.casefold()
    if suffix in {".xlsx", ".xlsm"}:
        return _xlsx_profile(path)
    if suffix in {".html", ".htm"}:
        return _html_table_profile(path)
    return _csv_profile(path)


def profile_for_source(path: Path, source_code: str) -> tuple[list[str], int, list[str], dict[str, Any]]:
    try:
        return profile_table(path)
    except ValueError:
        if source_code == "DJ" and path.suffix.casefold() in {".html", ".htm"}:
            return _dejan_html_profile(path)
        raise


def validate_office_archive(path: Path) -> None:
    try:
        with zipfile.ZipFile(path) as archive:
            members = archive.infolist()
            if len(members) > 10000:
                raise ValueError("Office workbook has too many ZIP members")
            if sum(item.file_size for item in members) > MAX_TABULAR_BYTES * 2:
                raise ValueError("Office workbook expands beyond the configured limit")
            if any(item.flag_bits & 0x1 for item in members):
                raise ValueError("Encrypted Office workbook members are unsupported")
            if any(item.file_size / max(1, item.compress_size) > 1000 for item in members):
                raise ValueError("Office workbook has a suspicious compression ratio")
    except zipfile.BadZipFile as exc:
        raise ValueError("Malformed Office workbook ZIP container") from exc


def file_contains_domain(path: Path, expected_domain: str) -> bool:
    needle = expected_domain.casefold()
    if path.suffix.casefold() in {".xlsx", ".xlsm"}:
        import openpyxl
        validate_office_archive(path)
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True, keep_links=False)
        try:
            return any(
                needle in str(value).casefold()
                for sheet in workbook.worksheets
                for row in sheet.iter_rows(values_only=True)
                for value in row if value is not None
            )
        finally:
            workbook.close()
    text, _ = decode_bytes(path.read_bytes())
    return needle in text.casefold()


def _file_snapshot(path: Path) -> str:
    return dt.datetime.fromtimestamp(path.stat().st_mtime, tz=dt.timezone.utc).replace(microsecond=0).isoformat().replace("+00:00", "Z")


def inspect_tabular(
    path: Path,
    source_code: str,
    expected_domain: str,
    freshness_class: str,
    explicit_binding: bool = True,
    valid_empty_attested: bool = False,
) -> dict[str, Any]:
    path = path.resolve()
    digest = sha256_file(path)
    try:
        headers, row_count, samples, transport = profile_for_source(path, source_code)
        mapping = header_mapping(headers)
        signature_ok, missing = signature_match(source_code, mapping)
        content_scope_match = file_contains_domain(path, expected_domain)
        contains_scope_column = bool({"target_url", "page"} & set(mapping)) and row_count > 0
        if contains_scope_column:
            scope_status = "match" if content_scope_match else "mismatch"
            scope_method = "content_match" if content_scope_match else "content_contradiction"
        else:
            scope_status = "match" if explicit_binding else "unknown"
            scope_method = "explicit_config_binding" if explicit_binding else "unproven"
        valid_empty = bool(row_count == 0 and valid_empty_attested and signature_ok)
        data_status = "present" if signature_ok and (row_count > 0 or valid_empty) else (
            "partial" if signature_ok else "malformed"
        )
        schema_payload = {
            "adapter_version": ADAPTER_VERSION,
            "source_code": source_code,
            "headers": [norm(header) for header in headers],
            "semantic_mapping": {key: norm(value) for key, value in sorted(mapping.items())},
            "structure": transport,
        }
        snapshot_at = _file_snapshot(path)
        schema_fingerprint = canonical_sha256(schema_payload)
        artifact = {
            "artifact_id": artifact_id(source_code, digest),
            "source_code": source_code,
            "path_or_locator": str(path),
            "sha256": digest,
            "byte_size": path.stat().st_size,
            "schema_fingerprint": schema_fingerprint,
            "snapshot_at": snapshot_at,
            "record_count": row_count,
            "valid_empty": valid_empty,
            "adapter_version": ADAPTER_VERSION,
        }
        freshness_status = freshness(snapshot_at, freshness_class)
        blocking = data_status != "present" or scope_status != "match" or freshness_status != "current"
        return {
            "artifact": artifact,
            "check": {
                "source_code": source_code,
                "access_status": "pass",
                "data_status": data_status,
                "scope_status": scope_status,
                "freshness_status": freshness_status,
                "blocking": blocking,
                "action": None if not blocking else (
                    f"Upload a current {source_code} export with the required schema and bind it explicitly to {expected_domain}."
                ),
                "details": {
                    "headers": headers,
                    "semantic_mapping": mapping,
                    "missing_semantics": missing,
                    "scope_method": scope_method,
                    **transport,
                },
            },
        }
    except Exception as exc:
        return {
            "artifact": {
                "artifact_id": artifact_id(source_code, digest),
                "source_code": source_code,
                "path_or_locator": str(path),
                "sha256": digest,
                "byte_size": path.stat().st_size,
                "schema_fingerprint": hashlib.sha256(b"malformed").hexdigest(),
                "snapshot_at": _file_snapshot(path),
                "record_count": 0,
                "valid_empty": False,
                "adapter_version": ADAPTER_VERSION,
            },
            "check": {
                "source_code": source_code,
                "access_status": "pass",
                "data_status": "malformed",
                "scope_status": "unknown",
                "freshness_status": freshness(_file_snapshot(path), freshness_class),
                "blocking": True,
                "action": f"Replace or re-export {source_code}; parser error: {type(exc).__name__}.",
                "details": {"error": str(exc)},
            },
        }


def parse_mhtml(path: Path) -> dict[str, Any]:
    path = path.resolve()
    if path.stat().st_size > MAX_MHTML_BYTES:
        raise ValueError(f"MHTML input exceeds {MAX_MHTML_BYTES} bytes")
    message = BytesParser(policy=policy.default).parsebytes(path.read_bytes())
    if not message.is_multipart() or message.get_content_subtype().casefold() != "related":
        raise ValueError("Expected MIME multipart/related MHTML")
    parts = []
    html_candidates: list[tuple[str, bytes]] = []
    total_decoded = 0
    for part in message.walk():
        if part.is_multipart():
            continue
        payload = part.get_payload(decode=True) or b""
        total_decoded += len(payload)
        if total_decoded > MAX_MHTML_BYTES * 2:
            raise ValueError("Decoded MHTML parts exceed the configured limit")
        location = str(part.get("Content-Location") or "")
        parts.append({
            "content_type": part.get_content_type(),
            "content_location": location,
            "content_id": str(part.get("Content-ID") or ""),
            "byte_size": len(payload),
            "sha256": hashlib.sha256(payload).hexdigest(),
        })
        if part.get_content_type() == "text/html":
            html_candidates.append((location, payload))
    if not html_candidates:
        raise ValueError("MHTML contains no text/html root")
    snapshot_location = str(message.get("Snapshot-Content-Location") or "")
    root_location, root_html = next(
        ((location, payload) for location, payload in html_candidates if snapshot_location and location == snapshot_location),
        html_candidates[0],
    )
    text, encoding = decode_bytes(root_html)
    parser = VisibleTextParser()
    parser.feed(text)
    visible_text = "\n".join(parser.text)
    title = " ".join(parser.title)
    return {
        "subject": str(message.get("Subject") or ""),
        "snapshot_content_location": snapshot_location or root_location or "",
        "root_content_location": root_location or "",
        "root_html_sha256": hashlib.sha256(root_html).hexdigest(),
        "root_encoding": encoding,
        "title": title,
        "visible_text": visible_text,
        "visible_text_sha256": hashlib.sha256(visible_text.encode("utf-8")).hexdigest(),
        "part_count": len(parts),
        "parts": parts,
        "remote_resources_loaded": False,
    }


def inspect_sentiment_mhtml(
    path: Path, expected_brand: str, expected_domain: str, freshness_class: str
) -> dict[str, Any]:
    path = path.resolve()
    digest = sha256_file(path)
    try:
        parsed = parse_mhtml(path)
        searchable = "\n".join([
            parsed["subject"], parsed["title"], parsed["snapshot_content_location"], parsed["visible_text"],
        ]).casefold()
        groups = [
            ("sentiment",),
            ("sentiment-score", "sentiment score", "tonalitat", "tonalität", "tonality"),
            ("lob", "praise", "positiv", "positive"),
            ("kritik", "criticism", "negativ", "negative"),
            ("chatgpt", "ai overview", "ai mode", "perplexity", "gemini", "copilot"),
        ]
        signature_hits = [any(term in searchable for term in group) for group in groups]
        signature_ok = sum(signature_hits) >= 4
        brand_match = expected_brand.casefold() in searchable
        domain_match = expected_domain.casefold() in searchable
        scope_status = "match" if brand_match and domain_match else "mismatch"
        snapshot_at = _file_snapshot(path)
        schema_fingerprint = canonical_sha256({
            "adapter_version": ADAPTER_VERSION,
            "content_type": "multipart/related",
            "signature_hits": signature_hits,
            "root_location_host_bound": bool(parsed["root_content_location"]),
        })
        data_status = "present" if signature_ok else "malformed"
        freshness_status = freshness(snapshot_at, freshness_class)
        blocking = data_status != "present" or scope_status != "match" or freshness_status != "current"
        return {
            "artifact": {
                "artifact_id": artifact_id("SX-SENT", digest),
                "source_code": "SX-SENT",
                "path_or_locator": str(path),
                "sha256": digest,
                "byte_size": path.stat().st_size,
                "schema_fingerprint": schema_fingerprint,
                "snapshot_at": snapshot_at,
                "record_count": 1,
                "valid_empty": False,
                "adapter_version": ADAPTER_VERSION,
            },
            "check": {
                "source_code": "SX-SENT",
                "access_status": "pass",
                "data_status": data_status,
                "scope_status": scope_status,
                "freshness_status": freshness_status,
                "blocking": blocking,
                "action": None if not blocking else "Upload a current SISTRIX Sentiment MHTML for the configured brand and domain.",
                "details": {
                    "brand_match": brand_match,
                    "domain_match": domain_match,
                    "signature_hits": signature_hits,
                    "subject": parsed["subject"],
                    "title": parsed["title"],
                    "snapshot_content_location": parsed["snapshot_content_location"],
                    "part_count": parsed["part_count"],
                    "remote_resources_loaded": False,
                    "visible_text_sha256": parsed["visible_text_sha256"],
                },
            },
            "parsed": parsed,
        }
    except Exception as exc:
        snapshot_at = _file_snapshot(path)
        return {
            "artifact": {
                "artifact_id": artifact_id("SX-SENT", digest),
                "source_code": "SX-SENT",
                "path_or_locator": str(path),
                "sha256": digest,
                "byte_size": path.stat().st_size,
                "schema_fingerprint": hashlib.sha256(b"malformed-mhtml").hexdigest(),
                "snapshot_at": snapshot_at,
                "record_count": 0,
                "valid_empty": False,
                "adapter_version": ADAPTER_VERSION,
            },
            "check": {
                "source_code": "SX-SENT",
                "access_status": "pass",
                "data_status": "malformed",
                "scope_status": "unknown",
                "freshness_status": freshness(snapshot_at, freshness_class),
                "blocking": True,
                "action": f"Save the complete SISTRIX Sentiment page as MHTML again; parser error: {type(exc).__name__}.",
                "details": {"error": str(exc), "remote_resources_loaded": False},
            },
        }
